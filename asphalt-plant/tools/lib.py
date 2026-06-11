"""Shared helpers for asphalt-plant doc builders.

Source of truth is YAML. Builders read YAML and emit md/xlsx/pdf so
the three formats can never drift.
"""
from __future__ import annotations

import datetime as _dt
from pathlib import Path
from typing import Any

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

REPO_ROOT = Path(__file__).resolve().parent.parent
TBD = "TBD"


def load_yaml(path: Path) -> dict[str, Any]:
    with path.open() as f:
        data = yaml.safe_load(f) or {}
    return data


def fmt(value: Any) -> str:
    """Render a YAML value for display. Empty/None -> TBD."""
    if value is None:
        return TBD
    if isinstance(value, bool):
        return "yes" if value else "no"
    s = str(value).strip()
    return s if s else TBD


def verified_tag(verified: Any) -> str:
    if verified is True:
        return "VERIFIED"
    if verified is False:
        return "ASSUMED"
    return TBD


# ---------- Excel helpers ----------

_HEADER_FILL = PatternFill("solid", fgColor="1F3864")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_THIN = Side(style="thin", color="888888")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def write_xlsx(path: Path, title: str, header_blocks: list[tuple[str, list[tuple[str, str]]]],
               table_title: str, columns: list[str], rows: list[list[Any]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31] or "Sheet1"

    r = 1
    ws.cell(row=r, column=1, value=title).font = Font(bold=True, size=14)
    r += 1
    ws.cell(row=r, column=1, value=f"generated {_dt.datetime.now().isoformat(timespec='seconds')}")
    r += 2

    for block_title, kvs in header_blocks:
        ws.cell(row=r, column=1, value=block_title).font = Font(bold=True)
        r += 1
        for k, v in kvs:
            ws.cell(row=r, column=1, value=k).font = Font(bold=True)
            ws.cell(row=r, column=2, value=v)
            r += 1
        r += 1

    ws.cell(row=r, column=1, value=table_title).font = Font(bold=True)
    r += 1

    for ci, col in enumerate(columns, start=1):
        c = ws.cell(row=r, column=ci, value=col)
        c.font = _HEADER_FONT
        c.fill = _HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = _BORDER
    r += 1

    for row in rows:
        for ci, val in enumerate(row, start=1):
            c = ws.cell(row=r, column=ci, value=val)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = _BORDER
        r += 1

    for ci, col in enumerate(columns, start=1):
        width = max(len(col), 12)
        for row in rows:
            cell_val = row[ci - 1] if ci - 1 < len(row) else ""
            width = max(width, min(40, len(str(cell_val)) + 1))
        ws.column_dimensions[get_column_letter(ci)].width = width

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# ---------- PDF helpers ----------

def write_pdf(path: Path, title: str, header_blocks: list[tuple[str, list[tuple[str, str]]]],
              table_title: str, columns: list[str], rows: list[list[Any]],
              landscape_page: bool = True) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    pagesize = landscape(LETTER) if landscape_page else LETTER
    doc = SimpleDocTemplate(
        str(path),
        pagesize=pagesize,
        leftMargin=0.4 * inch,
        rightMargin=0.4 * inch,
        topMargin=0.4 * inch,
        bottomMargin=0.4 * inch,
        title=title,
    )
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=styles["Heading1"], fontSize=14, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=11, spaceAfter=2)
    body = ParagraphStyle("body", parent=styles["BodyText"], fontSize=8, leading=10)
    small = ParagraphStyle("small", parent=styles["BodyText"], fontSize=7, textColor=colors.grey)

    story: list[Any] = [Paragraph(title, h1),
                        Paragraph(f"generated {_dt.datetime.now().isoformat(timespec='seconds')}", small),
                        Spacer(1, 6)]

    for block_title, kvs in header_blocks:
        story.append(Paragraph(block_title, h2))
        kv_rows = [[Paragraph(f"<b>{k}</b>", body), Paragraph(str(v), body)] for k, v in kvs]
        kv_tbl = Table(kv_rows, colWidths=[1.5 * inch, 5.5 * inch])
        kv_tbl.setStyle(TableStyle([
            ("BOX", (0, 0), (-1, -1), 0.25, colors.grey),
            ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.lightgrey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(kv_tbl)
        story.append(Spacer(1, 6))

    story.append(Paragraph(table_title, h2))
    table_data = [[Paragraph(f"<b>{c}</b>", body) for c in columns]]
    for row in rows:
        table_data.append([Paragraph(str(v) if v not in (None, "") else TBD, body) for v in row])

    table = Table(table_data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3864")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ALIGN", (0, 0), (-1, 0), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F2F2")]),
    ]))
    story.append(table)
    doc.build(story)


# ---------- Markdown helpers ----------

def md_kv_block(title: str, kvs: list[tuple[str, str]]) -> str:
    out = [f"### {title}", ""]
    for k, v in kvs:
        out.append(f"- **{k}:** {v}")
    out.append("")
    return "\n".join(out)


def md_table(columns: list[str], rows: list[list[Any]]) -> str:
    out = ["| " + " | ".join(columns) + " |",
           "|" + "|".join(["---"] * len(columns)) + "|"]
    for row in rows:
        cells = [str(v) if v not in (None, "") else TBD for v in row]
        out.append("| " + " | ".join(c.replace("|", "\\|").replace("\n", " ") for c in cells) + " |")
    out.append("")
    return "\n".join(out)


def write_md(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content)
